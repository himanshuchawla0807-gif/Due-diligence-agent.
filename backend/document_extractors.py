"""Shared document text extraction helpers."""

from __future__ import annotations

from pathlib import Path


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix == ".csv":
        return _extract_csv(path)
    if suffix == ".xlsx":
        return _extract_xlsx(path)
    return path.read_text(encoding="utf-8", errors="ignore")


def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception:
        try:
            import fitz

            with fitz.open(str(path)) as document:
                return "\n".join(page.get_text() or "" for page in document)
        except Exception:
            return ""


def _extract_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(part for part in parts if part)


def _extract_csv(path: Path) -> str:
    import pandas as pd

    return pd.read_csv(path).to_csv(index=False)


def _extract_xlsx(path: Path) -> str:
    import pandas as pd

    frames = pd.read_excel(path, sheet_name=None)
    return "\n\n".join(f"Sheet: {name}\n{frame.to_csv(index=False)}" for name, frame in frames.items())
