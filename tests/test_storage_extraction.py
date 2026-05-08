import shutil
from pathlib import Path

from backend.config import Settings
from backend.storage import SessionStore


def test_uploaded_documents_extract_docx_and_do_not_pretruncate():
    tmp_path = Path(".test-tmp/storage-extraction").resolve()
    if tmp_path.exists():
        shutil.rmtree(tmp_path, ignore_errors=True)

    try:
        source_dir = tmp_path / "source"
        source_dir.mkdir(parents=True)
        docx_path = source_dir / "management_report.docx"

        from docx import Document

        document = Document()
        document.add_paragraph("Revenue concentration risk in the enterprise customer base.")
        document.add_paragraph("Tail marker: " + ("working capital detail " * 2500))
        document.save(docx_path)

        settings = Settings(
            due_diligence_provider="mock",
            rag_storage_dir=tmp_path / "rag",
            session_store=tmp_path / "sessions.json",
            local_storage_dir=tmp_path / "storage",
            local_embedding_provider="local-hashing",
        )
        store = SessionStore(settings)
        session = store.get_or_create("docx-session")
        session["uploaded_files"].append(
            {
                "file_id": "docx-1",
                "filename": docx_path.name,
                "display_name": docx_path.name,
                "relative_path": "reports/management_report.docx",
                "path": str(docx_path),
                "size": docx_path.stat().st_size,
                "status": "stored",
            }
        )
        store.save()

        docs = store.documents("docx-session")

        assert docs[0]["file_name"] == "reports/management_report.docx"
        assert "Revenue concentration risk" in docs[0]["text"]
        assert len(docs[0]["text"]) > 24000
        assert "working capital detail" in docs[0]["text"]
    finally:
        shutil.rmtree(tmp_path, ignore_errors=True)
